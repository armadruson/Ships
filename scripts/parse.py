"""Normalise the NGO Shipbreaking Platform annual lists into one ship table.

Every yearly workbook has its own header layout (see README > Data), so columns
are matched by header name rather than by position.
"""
import csv
import datetime as dt
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "ships.csv"

FIELDS = ["year", "imo", "name", "type_raw", "category", "gt", "ldt", "built",
          "age", "flag", "owner_country", "place", "country", "arrival"]

# header text (upper-cased, whitespace-collapsed) -> field. Checked with startswith.
HEADER_RULES = [
    ("imo", ("IMO",)),
    ("name", ("NAME OF SHIP", "NAME")),
    ("type_raw", ("TYPE OF SHIP", "TYPE")),
    ("gt", ("GROSS TONNAGE", "GT")),
    ("ldt", ("LDT",)),
    ("built", ("BUILT",)),
    ("flag", ("LAST FLAG", "FLAG")),
    ("owner_country", ("BO COUNTRY", "BENEFICIAL OWNER'S COUNTRY", "COUNTRY OF THE BENEFICIAL")),
    ("place", ("PLACE", "DESTINATION CITY")),
    ("country", ("DESTINATION COUNTRY", "COUNTRY")),
    ("yard", ("DESTINATION YARD",)),
    ("arrival", ("ARRIVAL", "BEACHING DATE")),
]
# Headers that look like a field but are a different one.
HEADER_EXCLUDE = ("FLAG CHANGED", "FORMER", "PREVIOUS", "NEXT TO LAST", "FLAG PRIOR",
                  "CHANGE OF FLAG", "SUB-TYPE", "RO COUNTRY", "FORMER NAME")

# Order matters: first match wins.
CATEGORY_RULES = [
    ("Offshore & drilling", r"fpso|fso|drill|platform|rig\b|supply|anchor handling|pipe ?lay|diving|"
                            r"stand-?by|accommodation|seismo|well[- ]stim|crane|offshore|production testing|"
                            r"trenching|floating (production|storage|gas)|o\.r\.s\.v"),
    ("Fishing", r"fish|trawler"),
    ("Container", r"container ?ship|containership|fully cellular|barge container"),
    ("Tanker", r"tanker|gas carrier|lpg|lng|liquefied|oil carrier|\bobo\b|bulk and oil|bulk/oil|"
               r"bulk oil|ore/oil|ore and oil|ore/bulk/oil|bulk ore and oil"),
    ("Bulk carrier", r"bulk|bulker|bluk|ore carrier|wood[- ]?chips?|cement|limestone|aggregates|"
                     r"stone carrier|powder|sugar|ore trans"),
    ("Passenger & Ro-Ro", r"passeng|cruise|ferry|ro-?ro|roll on|vehicle"),
    ("General cargo", r"general cargo|genearl|cargo|reefer|refrigerated|livestock|livesstock|"
                      r"pallet|heavy[- ]load|pipe carrier|barge carrier|open hatch|semi-sub hl"),
]
CATEGORY_OTHER = "Other"

# Country spellings -> one canonical name.
COUNTRY_ALIASES = {
    "uk": "United Kingdom", "u.k.": "United Kingdom", "great britain": "United Kingdom",
    "usa": "United States", "us": "United States", "u.s.a.": "United States",
    "united states of america": "United States", "türkiye": "Turkey", "turkiye": "Turkey",
    "korea": "South Korea", "korea, south": "South Korea", "republic of korea": "South Korea",
    "s korea": "South Korea", "s. korea": "South Korea", "south korea": "South Korea",
    "uae": "United Arab Emirates", "u.a.e.": "United Arab Emirates", "u.a.e": "United Arab Emirates",
    "hong kong": "Hong Kong", "hong kong, china": "Hong Kong", "hongkong": "Hong Kong",
    "china (hong kong)": "Hong Kong", "st kitts & nevis": "St Kitts & Nevis",
    "saint kitts and nevis": "St Kitts & Nevis", "st. kitts & nevis": "St Kitts & Nevis",
    "st kitts and nevis": "St Kitts & Nevis", "antigua and barbuda": "Antigua & Barbuda",
    "st vincent & grenadines": "St Vincent & Grenadines", "st. vincent & grenadines": "St Vincent & Grenadines",
    "saint vincent and the grenadines": "St Vincent & Grenadines",
    "st vincent and the grenadines": "St Vincent & Grenadines",
    "marshall is": "Marshall Islands", "marshall is.": "Marshall Islands",
    "the netherlands": "Netherlands", "holland": "Netherlands", "russian federation": "Russia",
    "viet nam": "Vietnam", "tanzania, united republic of": "Tanzania", "cook islands": "Cook Islands",
    "sao tome & principe": "Sao Tome & Principe", "são tomé and príncipe": "Sao Tome & Principe",
    "sao tome and principe": "Sao Tome & Principe", "korea (south)": "South Korea",
    "korea (north)": "North Korea", "korea, north": "North Korea", "china, people's republic of": "China",
    "st kitts-nevis": "St Kitts & Nevis", "st. kitts-nevis": "St Kitts & Nevis",
    "st vincent & the grenadines": "St Vincent & Grenadines", "faeroe islands": "Faroe Islands",
    "irish republic": "Ireland", "tanzania (zanzibar)": "Tanzania", "spain (csr)": "Spain",
    "madeira": "Portugal", "hogn kong": "Hong Kong", "lybia": "Libya", "romenia": "Romania",
    "unknown": "", "unknwon": "", "n/a": "", "-": "", "?": "",
}
SMALL_WORDS = {"of", "and", "the", "&"}
# Flag registries carry a suffix for second/international registers: "Norway (Nis)" -> "Norway".
REGISTER_SUFFIX = re.compile(r"\s*\((nis|dis|mar|int.*|international.*|second.*)\)\s*$", re.I)

PLACE_ALIASES = {
    "chittagong": "Chattogram", "chattogram": "Chattogram", "alang": "Alang", "sosiya": "Alang",
    "alang-sosiya": "Alang", "gadani": "Gadani", "gadani beach": "Gadani", "aliaga": "Aliağa",
    "aliağa": "Aliağa", "jangjyin": "Jiangyin", "jiangyin,": "Jiangyin",
    "frederikshaven": "Frederikshavn", "sgravendeel": "Gravendeel", "'s-gravendeel": "Gravendeel", "unknown": "",
}


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def title(s):
    return " ".join(w if w.lower() in SMALL_WORDS and i else w.capitalize() if w.islower() or w.isupper() else w
                    for i, w in enumerate(s.split(" "))).replace(" Of ", " of ")


def norm_country(v):
    # "False" marks a fraudulent registration in the source; we keep the claimed flag.
    s = re.sub(r"\s+false$", "", REGISTER_SUFFIX.sub("", clean(v)), flags=re.I)
    key = s.lower()
    if key in COUNTRY_ALIASES:
        return COUNTRY_ALIASES[key]
    return title(s)


def norm_place(v):
    s = clean(v).strip(" ,")
    return PLACE_ALIASES.get(s.lower(), title(s))


def to_int(v):
    s = clean(v).replace(",", "")
    try:
        n = int(float(s))
        return n if n > 0 else None
    except ValueError:
        return None


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return ""


def categorise(type_raw):
    t = type_raw.lower()
    for cat, pattern in CATEGORY_RULES:
        if re.search(pattern, t):
            return cat
    return CATEGORY_OTHER


def map_header(row):
    """Return {field: column index} for a header row."""
    mapping = {}
    for idx, cell in enumerate(row):
        h = re.sub(r"\s+", " ", clean(cell).upper())
        if not h or h.startswith("*") or any(x in h for x in HEADER_EXCLUDE):
            continue
        for field, prefixes in HEADER_RULES:
            if field not in mapping and any(h.startswith(p) for p in prefixes):
                mapping[field] = idx
                break
    return mapping


def parse_workbook(path, year):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    # Header is row 1 (2012-14) or row 2 below a grouped super-header (2015+).
    head_idx, mapping = max(((i, map_header(r)) for i, r in enumerate(rows[:4])),
                            key=lambda x: len(x[1]))
    missing = {"imo", "name", "type_raw"} - mapping.keys()
    if missing:
        raise ValueError(f"{path.name}: could not find columns {missing}")

    def get(r, f):
        i = mapping.get(f)
        return r[i] if i is not None and i < len(r) else None

    ships = []
    for r in rows[head_idx + 1:]:
        imo = to_int(get(r, "imo"))
        name = clean(get(r, "name"))
        if not imo and not name:
            continue
        if "yard" in mapping:  # 2012/13: "Alang, India"
            yard = clean(get(r, "yard"))
            place, _, country = yard.rpartition(",") if "," in yard else ("", "", yard)
        else:
            place, country = get(r, "place"), clean(get(r, "country"))
            if "," in country:  # e.g. "Unknown shipbreakers, India"
                country = country.rpartition(",")[2]
        type_raw = clean(get(r, "type_raw"))
        built = to_int(get(r, "built"))
        built = built if built and 1900 < built <= year else None
        ships.append({
            "year": year,
            "imo": imo or "",
            "name": name.upper(),
            "type_raw": type_raw,
            "category": categorise(type_raw),
            "gt": to_int(get(r, "gt")) or "",
            "ldt": to_int(get(r, "ldt")) or "",
            "built": built or "",
            "age": (year - built) if built else "",
            "flag": norm_country(get(r, "flag")),
            "owner_country": norm_country(get(r, "owner_country")),
            "place": norm_place(place),
            "country": norm_country(country),
            "arrival": to_date(get(r, "arrival")),
        })
    return ships


def parse_all():
    ships = []
    for path in sorted(RAW.glob("*.xlsx")):
        ships += parse_workbook(path, int(path.stem))
    return ships


def main():
    ships = parse_all()
    with OUT.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(ships)
    print(f"parse: {len(ships)} ships -> {OUT.relative_to(ROOT)}")
    return ships


if __name__ == "__main__":
    main()
